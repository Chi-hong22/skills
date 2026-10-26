"""Input: 一份由 pandoc + assets/reference_cn.docx 生成的中文 docx；依赖 python-docx。
Output: 就地修好 pandoc 不处理的版式细节——表格撑满版心、压成三线表并加粗表头，题注与图片居中，时间范围行居中。
Pos: 中文报告版式管线的第二步（make_reference_cn.py 管样式，本脚本管单篇文档的局部版式）。
一旦本文件被更新，务必更新开头注释与 SKILL.md 的「中文报告版式」一节。

pandoc 写出的表格是 tblW=auto + 等分 gridCol，Word 会按等分列宽硬排，中文列名因此被折行；
这里改成 100% 版心宽 + autofit，让 Word 按内容分配列宽。

用法：
    python scripts/postprocess_cn_docx.py report.docx
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

CAPTION_RE = re.compile(r"^(表|图)\s*\d")
DRAWING = qn("w:drawing")


def set_border(parent, tag: str, val: str, size: int = 0) -> None:
    el = OxmlElement(tag)
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000" if val == "single" else "auto")
    parent.append(el)


def three_line(table) -> None:
    """三线表：表格只留上下两条粗线，表头行下加一条，其余行线去掉。"""
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for tag, val, size in (
        ("w:top", "single", 8),
        ("w:left", "none", 0),
        ("w:bottom", "single", 8),
        ("w:right", "none", 0),
        ("w:insideH", "none", 0),
        ("w:insideV", "none", 0),
    ):
        set_border(borders, tag, val, size)
    tbl_pr.append(borders)
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        old_tc = tc_pr.find(qn("w:tcBorders"))
        if old_tc is not None:
            tc_pr.remove(old_tc)
        tc_borders = OxmlElement("w:tcBorders")
        set_border(tc_borders, "w:bottom", "single", 8)
        tc_pr.append(tc_borders)


def fix_tables(doc) -> None:
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        width = tbl_pr.find(qn("w:tblW"))
        if width is None:
            width = OxmlElement("w:tblW")
            tbl_pr.append(width)
        width.set(qn("w:type"), "pct")
        width.set(qn("w:w"), "5000")
        layout = tbl_pr.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            tbl_pr.append(layout)
        layout.set(qn("w:type"), "autofit")
        if tbl_pr.find(qn("w:jc")) is None:
            jc = OxmlElement("w:jc")
            jc.set(qn("w:val"), "center")
            tbl_pr.append(jc)
        three_line(table)

        for index, row in enumerate(table.rows):
            if index == 0:
                tr_pr = row._tr.get_or_add_trPr()
                if tr_pr.find(qn("w:tblHeader")) is None:
                    tr_pr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                cell.vertical_alignment = 1  # 垂直居中
                for para in cell.paragraphs:
                    para.paragraph_format.first_line_indent = Pt(0)
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)
                    if index == 0:
                        for run in para.runs:
                            run.font.bold = True


def center(para) -> None:
    para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Pt(0)


def fix_paragraphs(doc) -> None:
    for para in doc.paragraphs:
        text = para.text.strip()
        if para._p.find(f".//{DRAWING}") is not None:
            center(para)
            continue
        if text.startswith("时间范围"):
            center(para)
            continue
        if CAPTION_RE.match(text):
            center(para)
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.keep_with_next = True
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10.5)


def main() -> int:
    parser = argparse.ArgumentParser(description="修补中文化 pandoc docx 的表格与题注版式")
    parser.add_argument("docx", type=Path, help="待修补的 docx（就地修改）")
    args = parser.parse_args()
    if not args.docx.is_file():
        print(f"not found: {args.docx}", file=sys.stderr)
        return 1
    doc = Document(str(args.docx))
    fix_tables(doc)
    fix_paragraphs(doc)
    doc.save(str(args.docx))
    print(f"patched: {args.docx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
